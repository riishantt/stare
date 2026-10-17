"""PDF / DOCX / TXT extraction (stare.intake.pdf_reader)."""

from __future__ import annotations

import io
from pathlib import Path

import pymupdf
import pytest

from stare.intake.pdf_reader import DocumentReadError, clean_text, extract_document

SAMPLES = Path(__file__).resolve().parents[2] / "data" / "samples"


def test_text_pdf_has_text_layer_and_no_ocr():
    doc = extract_document((SAMPLES / "complaint_ndtex_title7_retaliation.pdf").read_bytes(), "c.pdf")
    assert doc.page_count == 3 and len(doc.pages) == 3
    assert not doc.ocr_used and not any(p.ocr_used for p in doc.pages)
    assert "NORTHERN DISTRICT OF TEXAS" in doc.text
    assert "COUNT I" in doc.text and "Retaliation" in doc.text
    assert "Acme Logistics, Inc." in doc.text


def test_scanned_pdf_first_page_is_ocred():
    src = pymupdf.open(SAMPLES / "complaint_ndcal_ada_title7_scanned.pdf")
    first = pymupdf.open()
    first.insert_pdf(src, from_page=0, to_page=0)
    assert not first[0].get_text().strip()  # image-only: no text layer
    doc = extract_document(first.tobytes(), "scan.pdf")
    assert doc.ocr_used and doc.pages[0].ocr_used
    up = doc.text.upper()
    for word in ("NORTHERN DISTRICT", "CALIFORNIA", "PLAINTIFF", "BRIGHTHARBOR", "ACCOMMODATE"):
        assert word in up


def test_txt_and_docx_inputs():
    txt = extract_document(b"First line\n\nSecond   paragraph  here", "a.txt")
    assert txt.page_count == 1 and "Second paragraph here" in txt.text and "\n\n" in txt.text

    import docx

    d = docx.Document()
    d.add_paragraph("COMPLAINT")
    d.add_paragraph("Plaintiff alleges retaliation.")
    buf = io.BytesIO()
    d.save(buf)
    out = extract_document(buf.getvalue(), "x.docx")
    assert "Plaintiff alleges retaliation." in out.text and not out.ocr_used


def test_clean_text_dehyphenates_and_normalises():
    raw = "discrim-\nination   claims\r\nunder   Title VII\n\n\n\nNext   para"
    out = clean_text(raw)
    assert "discrimination claims" in out
    assert "\n\n\n" not in out and "  " not in out
    assert out.endswith("Next para")
    assert "Fifth-Circuit" in clean_text("Fifth-Circuit")  # real hyphens before capitals/words survive


@pytest.mark.parametrize(
    "data,name",
    [(b"", "e.pdf"), (b"not a pdf at all", "bad.pdf"), (b"MZ\x90\x00binary", "x.exe"), (b"PK\x03\x04junk", "b.docx"),
     (b"\x00\x01\x02\x03", "b.txt")],
)
def test_unsupported_or_corrupt_files_raise_clear_error(data, name):
    with pytest.raises(DocumentReadError):
        extract_document(data, name)
