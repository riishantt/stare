"""Agent and orchestrator tests with fake retrieval and a fake firewall (no corpus or network needed)."""

from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument

from stare.agents.base import AgentDeps
from stare.agents.firewall_agent import FirewallAgent
from stare.agents.memo import MemoAgent
from stare.agents.memo_firewall import check_memo, export_violations
from stare.agents.opposing import OpposingCounselAgent, procedural_flags
from stare.agents.precedent import PrecedentAgent, pick_quote
from stare.agents.research import ResearchAgent
from stare.agents.state import dump
from stare.db.models import CaseFile, Document, User
from stare.graph.events import EventBus
from stare.graph.service import ExportBlockedError, MemoVerificationError, ResearchService, RunStateError
from stare.llm import OfflineModel
from stare.schemas import (
    Authority,
    CaseRecord,
    CitationCheck,
    Claim,
    ExtractedCitation,
    GateResult,
    KeyDate,
    Memo,
    ResearchPlan,
    RetrievedChunk,
)

# --------------------------------------------------------------------------------------------- fakes

CASES = {
    "op:1": ("McDonnell Douglas Corp. v. Green", "411 U.S. 792", "scotus", 1973,
             "The complainant in a Title VII trial must carry the initial burden under the statute of establishing a prima facie case of discrimination."),
    "op:2": ("Univ. of Tex. Sw. Med. Ctr. v. Nassar", "570 U.S. 338", "scotus", 2013,
             "Title VII retaliation claims must be proved according to traditional principles of but-for causation."),
    "op:3": ("Burlington N. & S.F.R. Co. v. White", "548 U.S. 53", "scotus", 2006,
             "A plaintiff must show that a reasonable employee would have found the challenged action materially adverse."),
    "op:4": ("Smith v. Example Freight Co.", "999 F.3d 100", "ca9", 2020,
             "The employer must show that the plaintiff failed to establish protected activity under the opposition clause."),
    "op:5": ("Toyota Motor Mfg., Ky., Inc. v. Williams", "534 U.S. 184", "scotus", 2002,
             "The terms of the statute need to be interpreted strictly to create a demanding standard for qualifying as disabled."),
    "op:6": ("Clark County School Dist. v. Breeden", "532 U.S. 268", "scotus", 2001,
             "The temporal proximity must be very close when mere temporal proximity is offered as evidence of causality."),
}


def _authority(key: str, query: str) -> Authority:
    name, cite, court, year, text = CASES[key]
    oid = int(key.split(":")[1])
    chunk = RetrievedChunk(chunk_id=oid, kind="opinion", opinion_id=oid, text=text, query=query, rrf_score=0.5)
    return Authority(key=key, kind="case", citation=cite, case_name=name, court_id=court, year=year,
                     opinion_id=oid, relevance=0.9 - oid * 0.05, passages=[chunk])


class FakeRetriever:
    def __init__(self):
        self.queries: list[list[str]] = []

    def search_many(self, queries, k=30, **kw):
        self.queries.append(list(queries))
        adverse = any("temporal proximity alone" in q or "petty slights" in q for q in queries)
        keys = ["op:6", "op:4"] if adverse else ["op:1", "op:2", "op:3", "op:5"]
        return [c for key in keys for c in _authority(key, queries[0]).passages]

    def authorities_from_chunks(self, chunks, limit):
        out, seen = [], set()
        for c in chunks:
            key = f"op:{c.opinion_id}"
            if key not in seen:
                seen.add(key)
                out.append(_authority(key, c.query))
        return out[:limit]


class FakeFirewall:
    """Verifies the known cases; Toyota gets a Gate 4 warning; anything else is blocked."""

    def _check(self, cite: str, name: str = "") -> CitationCheck:
        known = {v[1]: v[0] for v in CASES.values()}
        gates = [GateResult(gate=1, name="extract", outcome="pass")]
        if cite not in known:
            gates.append(GateResult(gate=2, name="exists", outcome="fail"))
            return CitationCheck(citation=ExtractedCitation(raw=f"{name}, {cite}", cite=cite, case_name=name),
                                 gates=gates, status="blocked", reason="no matching opinion in index or API")
        gates.append(GateResult(gate=2, name="exists", outcome="pass"))
        if cite == "534 U.S. 184":
            gates.append(GateResult(gate=4, name="good_law", outcome="warn"))
            return CitationCheck(citation=ExtractedCitation(raw=cite, cite=cite), gates=gates, status="warning",
                                 reason="superseded by statute", treatment=["ADA Amendments Act of 2008"])
        return CitationCheck(citation=ExtractedCitation(raw=cite, cite=cite), gates=gates, status="verified",
                             reason="quote matched; no negative treatment")

    def verify_authority(self, a: Authority) -> CitationCheck:
        return self._check(a.citation, a.case_name)

    def verify_text(self, text: str) -> list[CitationCheck]:
        import re

        return [self._check(m.group(1)) for m in re.finditer(r"(\d+ (?:U\.S\.|F\.3d) \d+)", text)]

    def filter_text(self, text: str, checks):
        bad = [c.citation.cite for c in checks if not c.exportable]
        sentences = [s for s in text.split(". ") if not any(b in s for b in bad)]
        return ". ".join(sentences), bad


@pytest.fixture
def deps(session_factory, settings):
    d = AgentDeps(llm=OfflineModel(), session_factory=session_factory, settings=settings)
    d.__dict__["retriever"] = FakeRetriever()
    d.__dict__["firewall"] = FakeFirewall()
    return d


RECORD = CaseRecord(
    court_id="txnd", court_name="N.D. Tex.", circuit="5",
    claims=[Claim(statute="title_vii", theory="retaliation")],
    key_dates=[KeyDate(label="Termination", date="2025-07-15"), KeyDate(label="EEOC charge filed", date="2026-06-01")],
    facts_summary="Plaintiff complained to human resources about race discrimination and was terminated two weeks later.",
)


def _base_state() -> dict:
    from stare.intake.planner import draft_plan

    return {"run_id": 1, "case_record": dump(RECORD), "plan": dump(draft_plan(RECORD)),
            "redacted_text": "Plaintiff complained to human resources about discrimination. Plaintiff was terminated after the complaint."}


# --------------------------------------------------------------------------------------------- units


def test_procedural_flags_detect_late_charge():
    flags = procedural_flags(RECORD)
    assert any("beyond the 300-day limit" in f for f in flags)


def test_pick_quote_prefers_rule_sentence():
    a = _authority("op:2", "but-for causation")
    assert "but-for causation" in pick_quote(a, "but-for causation retaliation")


def test_research_then_precedent_ranking(deps):
    state = _base_state()
    state |= ResearchAgent(deps).run(state)
    assert state["candidates"] and all(c["plan_item_id"] for c in state["candidates"])
    assert any("but-for" in q for qs in deps.retriever.queries for q in qs)  # element queries were used
    state |= PrecedentAgent(deps).run(state)
    ranks = {a["citation"]: a["rank"] for a in state["shortlist"]}
    assert ranks["411 U.S. 792"]["status"] == "binding"
    assert ranks["999 F.3d 100"]["status"] == "persuasive" if "999 F.3d 100" in ranks else True
    assert state["timings"]["research"] >= 0


def test_firewall_agent_flags_superseded_case_and_memo_lists_warning(deps):
    state = _base_state()
    for agent in (ResearchAgent, PrecedentAgent, OpposingCounselAgent, FirewallAgent):
        state |= agent(deps).run(state)
    toyota = next(a for a in state["verified"] if a["citation"] == "534 U.S. 184")
    assert toyota["rank"]["status"] == "flagged"
    assert state["adverse"] and state["adverse_report"]
    state |= MemoAgent(deps).run(state)
    memo = Memo.model_validate(state["memo"])
    assert "534 U.S. 184" not in memo.rule  # flagged authority is not used as support
    assert any("Toyota" in w for w in memo.warnings)
    assert "Breeden" in memo.adverse_authority
    assert memo.citations and all(c.marker in (memo.rule + memo.adverse_authority) for c in memo.citations)


def test_check_memo_removes_hallucinated_citation(deps):
    memo = Memo(title="t", question_presented="q", issue="i",
                rule="Retaliation requires but-for causation, Univ. of Tex. Sw. Med. Ctr. v. Nassar, 570 U.S. 338 (2013) [1]. "
                     "Airlines are liable, Varghese v. China Southern Airlines, 925 F.3d 1339 (11th Cir. 2019) [2]",
                application="a", conclusion="c")
    cleaned, checks = check_memo(memo, deps.firewall)
    assert "925 F.3d 1339" not in cleaned.rule and "570 U.S. 338" in cleaned.rule
    assert any("925 F.3d 1339" in r for r in cleaned.removed_citations)
    assert export_violations(cleaned, deps.firewall) == []
    assert export_violations(memo, deps.firewall)


# --------------------------------------------------------------------------------------------- graph


@pytest.fixture
def service(session_factory, settings, deps, monkeypatch):
    svc = ResearchService(session_factory, settings, OfflineModel(), synchronous=True, event_bus=EventBus())
    svc.deps.__dict__["retriever"] = deps.retriever
    svc.deps.__dict__["firewall"] = deps.firewall
    import stare.agents.intake as intake_mod

    monkeypatch.setattr(intake_mod, "extract_case_record", lambda text, llm, h: RECORD)
    with session_factory() as s:
        attorney = User(name="Daniel Ortiz", email="a@x", password_hash="x", role="attorney")
        s.add(attorney)
        s.flush()
        case = CaseFile(matter_number="M-0001", title="[PLAINTIFF] v. Acme", user_id=attorney.id)
        s.add(case)
        s.flush()
        s.add(Document(case_id=case.id, file_name="c.pdf", redacted_text="[PLAINTIFF] complained and was fired."))
        s.commit()
        svc._ids = (attorney.id, case.id)
    return svc


def test_full_run_with_approval_gates_and_export(service):
    user_id, case_id = service._ids
    sub = service.subscribe(1)
    run_id = service.start_run(case_id, user_id)
    snap = service.snapshot(run_id)
    assert snap["status"] == "awaiting_plan_approval" and snap["plan"]["items"]
    with pytest.raises(RunStateError):  # memo gate cannot be skipped
        service.approve_memo(run_id, None, user_id)

    service.approve_plan(run_id, ResearchPlan.model_validate(snap["plan"]), user_id)
    snap = service.snapshot(run_id)
    assert snap["status"] == "awaiting_memo_approval"
    assert snap["memo"] and snap["latency_ms"] is not None
    agents = [s["agent"] for s in snap["steps"] if s["status"] == "done"]
    assert agents[:6] == ["intake", "research", "precedent", "opposing_counsel", "citation_firewall", "memo"]
    with pytest.raises(RunStateError):  # export requires approval
        service.export_docx(run_id, user_id)

    service.request_memo_changes(run_id, "Add more on causation.", user_id)
    snap = service.snapshot(run_id)
    assert snap["status"] == "awaiting_memo_approval" and snap["memo_version"] == 2

    bad = Memo.model_validate(snap["memo"])
    bad.rule += " See Varghese v. China Southern Airlines, 925 F.3d 1339 (11th Cir. 2019)."
    with pytest.raises(MemoVerificationError):
        service.approve_memo(run_id, bad, user_id)

    service.approve_memo(run_id, None, user_id)
    assert service.snapshot(run_id)["status"] == "approved"
    data = service.export_docx(run_id, user_id)
    text = "\n".join(p.text for p in DocxDocument(io.BytesIO(data)).paragraphs)
    assert "Rule" in text and "Adverse Authority" in text
    assert service.snapshot(run_id)["status"] == "exported"
    events = []
    while (e := sub.get(timeout=0.01)) is not None:
        events.append(e)
    assert any(e["type"] == "trace" for e in events) and any(e.get("status") == "exported" for e in events)


def test_export_blocked_when_memo_contains_unverified_citation(service, session_factory):
    from stare.db.models import MemoRow

    user_id, case_id = service._ids
    run_id = service.start_run(case_id, user_id)
    service.approve_plan(run_id, ResearchPlan.model_validate(service.snapshot(run_id)["plan"]), user_id)
    service.approve_memo(run_id, None, user_id)
    with session_factory() as s:  # simulate tampering with the stored memo after approval
        m = s.query(MemoRow).filter_by(run_id=run_id).order_by(MemoRow.version.desc()).first()
        m.irac_json = {**m.irac_json, "rule": "See Miller v. United Airlines, Inc., 174 F.3d 366 (4th Cir. 1999)."}
        s.commit()
    with pytest.raises(ExportBlockedError):
        service.export_docx(run_id, user_id)


def test_stale_plan_approval_cannot_resume_the_memo_gate(service, monkeypatch):
    """Two approvals racing past the status check: the second must not reach the memo interrupt."""
    user_id, case_id = service._ids
    run_id = service.start_run(case_id, user_id)
    plan = ResearchPlan.model_validate(service.snapshot(run_id)["plan"])
    stale = service._require(run_id, "awaiting_plan_approval")
    service.approve_plan(run_id, plan, user_id)
    assert service.snapshot(run_id)["status"] == "awaiting_memo_approval"

    monkeypatch.setattr(service, "_require", lambda rid, status: stale)  # the second request read the old status
    with pytest.raises(RunStateError):
        service.approve_plan(run_id, plan, user_id)
    snap = service.snapshot(run_id)
    assert snap["status"] == "awaiting_memo_approval" and snap["memo_approved_by"] is None


def test_edited_memo_cannot_inject_references_or_unchecked_citations(service):
    from stare.schemas import MemoCitation

    user_id, case_id = service._ids
    run_id = service.start_run(case_id, user_id)
    service.approve_plan(run_id, ResearchPlan.model_validate(service.snapshot(run_id)["plan"]), user_id)
    original = Memo.model_validate(service.snapshot(run_id)["memo"])

    bad = original.model_copy(deep=True)
    bad.question_presented += " See Varghese v. China Southern Airlines, 925 F.3d 1339 (11th Cir. 2019)."
    with pytest.raises(MemoVerificationError):
        service.approve_memo(run_id, bad, user_id)

    forged = original.model_copy(deep=True)
    forged.conclusion += " [99]"
    forged.citations.append(MemoCitation(marker="[99]", citation="925 F.3d 1339", case_name="Varghese", status="verified"))
    forged.warnings, forged.title = [], "Forged"
    service.approve_memo(run_id, forged, user_id)
    stored = Memo.model_validate(service.snapshot(run_id)["memo"])
    assert stored.conclusion.endswith("[99]")  # the prose edit is kept
    body = f"{stored.question_presented} {stored.issue} {stored.rule} {stored.application} {stored.conclusion} {stored.adverse_authority}"
    assert [c.marker for c in stored.citations] == [c.marker for c in original.citations if c.marker in body]
    assert "925 F.3d 1339" not in [c.citation for c in stored.citations]
    assert stored.title == original.title and stored.warnings == original.warnings


def test_memo_warning_lists_gate4_treatment_once():
    from stare.agents.memo import _warning_line
    from stare.schemas import Authority, CitationCheck, ExtractedCitation

    t = "superseded by statute: ADA Amendments Act of 2008"
    check = CitationCheck(citation=ExtractedCitation(raw="534 U.S. 184", cite="534 U.S. 184"), status="warning",
                          reason=f"good_law: {t}", treatment=[t, "called into doubt by Later v. Case"])
    a = Authority(key="op:1", kind="case", citation="534 U.S. 184", case_name="Toyota v. Williams", verification=check)
    line = _warning_line(a)
    assert line.count(t) == 1 and line.endswith("(called into doubt by Later v. Case)")


def test_quoted_rule_sentence_keeps_weight_with_its_own_citation():
    from stare.citations.parser import extract_citations, sentence_starts

    text = ('In A v. B, 1 U.S. 2 (1990) [1] (Supreme Court; binds every federal court), the court explained: '
            '"The first rule is stated here." In C v. D, 3 U.S. 4 (1991) [2] (Supreme Court; binds every federal court), '
            'the court explained: "The second rule is stated here."')
    cites = extract_citations(text)
    assert [c.quote for c in cites] == ["The first rule is stated here.", "The second rule is stated here."]
    starts = sentence_starts(text, [c.span for c in cites])
    assert text.index("In C v. D") in starts and text.index("(Supreme Court") not in starts
