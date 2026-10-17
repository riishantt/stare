"""LLM mode (STARE_LLM_PROVIDER=anthropic / openai_compat) with scripted fake models: never a real API.

The rest of the suite runs the deterministic offline path. These tests drive the same agents and graph
through their LLM branches, and check the provider adapters against fake SDK clients.
"""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from test_agents import RECORD, FakeFirewall, FakeRetriever

from stare.agents.memo import _Draft
from stare.agents.opposing import _Counter, _Reason
from stare.agents.research import _Queries
from stare.db.models import CaseFile, Document, User
from stare.graph.events import EventBus
from stare.graph.service import ResearchService
from stare.intake.planner import _PlanDraft
from stare.llm import get_llm
from stare.llm.base import ChatModel, LLMError, UnredactedTextError
from stare.schemas import Memo, ResearchPlan

FAKE_CITE = "Varghese v. China Southern Airlines, 925 F.3d 1339 (11th Cir. 2019)"


class ScriptedLLM(ChatModel):
    """Answers each agent's structured request; the memo draft smuggles in a fabricated Mata citation."""

    provider = "scripted"
    is_offline = False

    def __init__(self, fail: bool = False):
        super().__init__("scripted-main", "scripted-fast")
        self.fail = fail
        self.calls: list[tuple[str, str]] = []

    def _complete(self, system, prompt, *, tier, max_tokens):
        raise LLMError("unused")

    def _complete_json(self, system, prompt, schema, *, tier, max_tokens):
        self.calls.append((schema.__name__, prompt))
        if self.fail:
            raise LLMError("model unavailable")
        if schema is _PlanDraft:
            return _PlanDraft(question="LLM-refined question: does retaliation causation hold in N.D. Tex.?",
                              items=[{"id": "P1", "issue": "Elements of a Title VII retaliation claim", "statute": "title_vii",
                                      "elements": ["protected activity", "materially adverse action", "but-for causation"]}])
        if schema is _Queries:
            return _Queries(queries=["llm query: but-for causation retaliation Fifth Circuit"])
        if schema is _Counter:
            return _Counter(queries=["llm counter: temporal proximity alone insufficient"])
        if schema is _Reason:
            return _Reason(reason="LLM reason: the defence will argue the gap breaks causation.")
        if schema is _Draft:
            return _Draft(
                issue="Whether the termination was retaliation.",
                rule="Retaliation requires but-for causation, Univ. of Tex. Sw. Med. Ctr. v. Nassar, 570 U.S. 338 (2013) [1]. "
                     f"Airlines are always liable, {FAKE_CITE} [9]",
                application="The two-week gap supports causation.",
                conclusion="The claim is viable, subject to attorney review.",
            )
        raise AssertionError(f"unexpected schema {schema.__name__}")


@pytest.fixture
def make_service(session_factory, settings, monkeypatch):
    import stare.agents.intake as intake_mod

    monkeypatch.setattr(intake_mod, "extract_case_record", lambda text, llm, h: RECORD)

    def build(llm: ChatModel) -> ResearchService:
        svc = ResearchService(session_factory, settings, llm, synchronous=True, event_bus=EventBus())
        svc.deps.__dict__["retriever"] = FakeRetriever()
        svc.deps.__dict__["firewall"] = FakeFirewall()
        with session_factory() as s:
            attorney = User(name="Aisha Brooks", email=f"a{id(svc)}@x", password_hash="x", role="attorney")
            s.add(attorney)
            s.flush()
            case = CaseFile(matter_number=f"M-{id(svc)}", title="[PLAINTIFF] v. Acme", user_id=attorney.id)
            s.add(case)
            s.flush()
            s.add(Document(case_id=case.id, file_name="c.pdf", redacted_text="[PLAINTIFF] complained and was fired."))
            s.commit()
            svc._ids = (attorney.id, case.id)
        return svc

    return build


def test_llm_mode_full_run_firewall_strips_fabricated_citation(make_service):
    llm = ScriptedLLM()
    svc = make_service(llm)
    user_id, case_id = svc._ids
    run_id = svc.start_run(case_id, user_id)
    snap = svc.snapshot(run_id)
    assert snap["plan"]["question"].startswith("LLM-refined")  # planner used the model

    svc.approve_plan(run_id, ResearchPlan.model_validate(snap["plan"]), user_id)
    snap = svc.snapshot(run_id)
    assert snap["status"] == "awaiting_memo_approval"
    memo = Memo.model_validate(snap["memo"])
    assert memo.generated_by == "scripted:scripted-main"
    assert "570 U.S. 338" in memo.rule
    assert "925 F.3d 1339" not in memo.rule  # the memo Firewall pass removed the fabricated case
    assert any("925 F.3d 1339" in r for r in memo.removed_citations)
    assert "LLM reason" in memo.adverse_authority

    used = {name for name, _ in llm.calls}
    assert {"_PlanDraft", "_Queries", "_Counter", "_Reason", "_Draft"} <= used
    assert all("Rivera" not in prompt for _, prompt in llm.calls)  # only redacted text reaches the model

    svc.approve_memo(run_id, None, user_id)
    lines = [p.text for p in DocxDocument(io.BytesIO(svc.export_docx(run_id, user_id))).paragraphs]
    assert any("570 U.S. 338" in line for line in lines)
    # the fabricated case appears only in the list of citations the Firewall removed
    assert [line for line in lines if "925 F.3d 1339" in line] == [
        "925 F.3d 1339: removed from the memo by the Citation Firewall"]


def test_llm_failures_fall_back_to_the_offline_path(make_service):
    llm = ScriptedLLM(fail=True)
    svc = make_service(llm)
    user_id, case_id = svc._ids
    run_id = svc.start_run(case_id, user_id)
    svc.approve_plan(run_id, ResearchPlan.model_validate(svc.snapshot(run_id)["plan"]), user_id)
    snap = svc.snapshot(run_id)
    assert snap["status"] == "awaiting_memo_approval" and llm.calls
    memo = Memo.model_validate(snap["memo"])
    assert memo.rule  # extractive draft replaced the failed LLM draft
    assert memo.generated_by == "offline fallback (scripted:scripted-main unavailable)"  # never reported as LLM-drafted
    assert llm.usage.failures == len(llm.calls) and llm.usage.calls == 0


# -- provider plumbing --------------------------------------------------------------------------------


class _Echo(ChatModel):
    provider = "echo"

    def __init__(self, cache_dir=None):
        super().__init__("m", "f", cache_dir)
        self.sent = 0

    def _complete(self, system, prompt, *, tier, max_tokens):
        self.sent += 1
        return f"reply {self.sent}"

    def _complete_json(self, system, prompt, schema, *, tier, max_tokens):
        self.sent += 1
        return schema(queries=[prompt])


@pytest.mark.parametrize("pii", ["SSN 123-45-6789", "mail jane.doe@example.com", "call (214) 555-0182"])
def test_pii_guard_refuses_before_any_call(pii):
    llm = _Echo()
    with pytest.raises(UnredactedTextError):
        llm.complete("s", f"Facts: {pii}")
    with pytest.raises(UnredactedTextError):
        llm.complete_json("s", f"Facts: {pii}", _Queries)
    assert llm.sent == 0


def test_disk_cache_avoids_repeat_calls(tmp_path):
    llm = _Echo(cache_dir=tmp_path)
    assert llm.complete("s", "p") == llm.complete("s", "p") == "reply 1"
    assert llm.complete_json("s", "q", _Queries).queries == llm.complete_json("s", "q", _Queries).queries == ["q"]
    assert llm.sent == 2
    assert llm.complete("s", "p", tier="fast") == "reply 3"  # different model tier, different cache entry


def test_factory_selects_provider(settings, monkeypatch):
    assert get_llm(settings).is_offline
    monkeypatch.setattr(settings, "llm_provider", "openai_compat")
    monkeypatch.setattr(settings, "llm_base_url", "http://localhost:11434/v1")
    llm = get_llm(settings)
    assert llm.provider == "openai_compat" and not llm.is_offline
    monkeypatch.setattr(settings, "llm_base_url", "")
    with pytest.raises(LLMError):
        get_llm(settings)


def _openai_reply(content: str):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))],
                           usage=SimpleNamespace(prompt_tokens=11, completion_tokens=7))


def test_openai_compat_parses_json_and_wraps_errors():
    from stare.llm.openai_compat import OpenAICompatModel

    llm = OpenAICompatModel("qwen", "qwen-small", "http://localhost:11434/v1")
    seen = {}

    def create(**kw):
        seen.update(kw)
        return _openai_reply('Sure:\n```json\n{"queries": ["a", "b"]}\n```')

    llm.client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    assert llm.complete_json("sys", "p", _Queries, tier="fast").queries == ["a", "b"]
    assert seen["model"] == "qwen-small" and seen["response_format"] == {"type": "json_object"}
    assert "JSON Schema" in seen["messages"][0]["content"]
    assert llm.usage.calls == 1 and llm.usage.cost_usd == 0.0

    llm.client.chat.completions.create = lambda **kw: _openai_reply("not json")
    with pytest.raises(LLMError, match="invalid JSON"):
        llm.complete_json("sys", "p2", _Queries)

    def boom(**kw):
        raise ConnectionError("refused")

    llm.client.chat.completions.create = boom
    with pytest.raises(LLMError, match="model server error"):
        llm.complete("sys", "p3")


def _claude_reply(text: str, stop: str = "end_turn"):
    return SimpleNamespace(stop_reason=stop, content=[SimpleNamespace(type="text", text=text)],
                           usage=SimpleNamespace(input_tokens=1_000_000, output_tokens=100_000))


def test_anthropic_adapter_cost_schema_and_stop_reasons():
    pytest.importorskip("anthropic")
    from stare.llm.anthropic_llm import AnthropicModel

    llm = AnthropicModel("claude-opus-5-5", "claude-haiku-5-5", api_key="test-key-never-used")
    sent = {}

    def create(**kw):
        sent.update(kw)
        return _claude_reply('{"queries": ["x"]}')

    llm.client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=create)))
    assert llm.complete_json("sys", "p", _Queries, tier="fast").queries == ["x"]
    assert sent["model"] == "claude-haiku-5-5"
    assert sent["output_config"]["format"]["schema"]["additionalProperties"] is False
    assert llm.usage.cost_usd == pytest.approx(0.10 + 0.05)

    llm.client.beta.messages.create = lambda **kw: _claude_reply("", stop="refusal")
    with pytest.raises(LLMError, match="declined"):
        llm.complete("sys", "p2")
    llm.client.beta.messages.create = lambda **kw: _claude_reply("partial", stop="max_tokens")
    with pytest.raises(LLMError, match="truncated"):
        llm.complete("sys", "p3")
