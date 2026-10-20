"""HTTP API tests: TestClient against a synchronous ResearchService with fake retriever/firewall."""

from __future__ import annotations

import io
import json

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from stare.api.auth import create_user
from stare.api.main import create_app
from stare.api.routers import misc
from stare.citations.firewall import CitationFirewall
from stare.db.models import AuditLog, CaseFile, CitationEdge, Document, Opinion
from stare.graph.events import EventBus
from stare.graph.service import ResearchService
from stare.intake.redactor import PIIRedactor
from stare.llm import OfflineModel
from test_agents import RECORD, FakeFirewall, FakeRetriever

PW = "correct-horse-battery"


@pytest.fixture
def service(session_factory, settings, monkeypatch):
    settings.jwt_secret = "test-secret-" + "x" * 32
    svc = ResearchService(session_factory, settings, OfflineModel(), synchronous=True, event_bus=EventBus())
    svc.deps.__dict__["retriever"] = FakeRetriever()
    svc.deps.__dict__["firewall"] = FakeFirewall()
    import stare.agents.intake as intake_mod

    monkeypatch.setattr(intake_mod, "extract_case_record", lambda text, llm, h: RECORD)
    monkeypatch.setattr(misc, "_redactor", PIIRedactor("regex"))  # keep tests light: no Presidio/spaCy
    return svc


@pytest.fixture
def client(service, session_factory):
    with session_factory() as s:
        create_user(s, "Admin", "admin@firm.test", PW, "admin")
        create_user(s, "Daniel Ortiz", "atty@firm.test", PW, "attorney")
        create_user(s, "Paula Legal", "para@firm.test", PW, "paralegal")
        s.commit()
    with TestClient(create_app(service)) as c:
        yield c


def login(client, who: str) -> dict:
    r = client.post("/api/auth/login", json={"email": f"{who}@firm.test", "password": PW})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture
def atty(client):
    return login(client, "atty")


@pytest.fixture
def para(client):
    return login(client, "para")


@pytest.fixture
def admin(client):
    return login(client, "admin")


def make_matter(client, headers) -> int:
    r = client.post("/api/matters", json={"title": "[PLAINTIFF] v. Acme", "court_id": "txnd"}, headers=headers)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def upload(client, headers, matter_id, text="Jane complained about bias. SSN 123-45-6789, mail jane.doe@example.com."):
    return client.post(f"/api/matters/{matter_id}/documents", headers=headers,
                       files={"file": ("intake.txt", text.encode(), "text/plain")})


# ----------------------------------------------------------------------------------------------- auth


def test_health_needs_no_auth(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "ok" and body["llm_provider"] == "offline"
    assert set(body["corpus"]) == {"opinions", "statutes", "chunks"} and body["index_built"] is False


def test_login_me_and_failures(client, session_factory):
    assert client.get("/api/auth/me").status_code == 401
    assert client.post("/api/auth/login", json={"email": "atty@firm.test", "password": "nope"}).status_code == 401
    assert client.post("/api/auth/login", json={"email": "ghost@firm.test", "password": PW}).status_code == 401
    h = login(client, "atty")
    me = client.get("/api/auth/me", headers=h).json()
    assert me["role"] == "attorney" and "password" not in json.dumps(me)
    assert client.get("/api/auth/me", headers={"Authorization": "Bearer garbage"}).status_code == 401
    with session_factory() as s:
        rows = s.query(AuditLog).filter(AuditLog.action.like("auth.%")).all()
        assert {r.action for r in rows} == {"auth.login", "auth.login_failed"}
        assert "nope" not in json.dumps([r.detail for r in rows])


def test_inactive_user_rejected(client, admin):
    users = client.get("/api/users", headers=admin).json()
    para_id = next(u["id"] for u in users if u["role"] == "paralegal")
    h = login(client, "para")
    assert client.patch(f"/api/users/{para_id}", json={"active": False}, headers=admin).status_code == 200
    assert client.get("/api/auth/me", headers=h).status_code == 401  # existing token stops working
    assert client.post("/api/auth/login", json={"email": "para@firm.test", "password": PW}).status_code == 401


def test_users_admin_crud_and_roles(client, admin, atty):
    assert client.get("/api/users", headers=atty).status_code == 403
    assert client.get("/api/users").status_code == 401
    r = client.post("/api/users", headers=admin,
                    json={"name": "New", "email": "New@Firm.test", "password": "longenough1", "role": "paralegal"})
    assert r.status_code == 201 and r.json()["email"] == "new@firm.test"
    uid = r.json()["id"]
    assert client.post("/api/users", headers=admin, json={"name": "Dup", "email": "new@firm.test",
                                                          "password": "longenough1", "role": "attorney"}).status_code == 409
    r = client.patch(f"/api/users/{uid}", headers=admin, json={"role": "attorney", "password": "anotherpass9"})
    assert r.json()["role"] == "attorney"
    assert client.post("/api/auth/login", json={"email": "new@firm.test", "password": "anotherpass9"}).status_code == 200
    me_id = next(u["id"] for u in client.get("/api/users", headers=admin).json() if u["role"] == "admin")
    assert client.patch(f"/api/users/{me_id}", headers=admin, json={"active": False}).status_code == 409
    assert client.patch("/api/users/9999", headers=admin, json={"active": False}).status_code == 404


def test_courts_list(client, atty):
    courts = client.get("/api/courts", headers=atty).json()
    assert any(c["id"] == "txnd" for c in courts)


# ----------------------------------------------------------------------------------------------- matters


def test_matter_numbers_and_listing(client, atty, para):
    a = client.post("/api/matters", json={"title": "A v. B"}, headers=atty).json()
    b = client.post("/api/matters", json={"title": "C v. D", "court_id": "ca5"}, headers=para).json()
    assert (a["matter_number"], b["matter_number"]) == ("M-0001", "M-0002")
    assert b["court_name"]
    assert client.post("/api/matters", json={"title": "X", "court_id": "nope"}, headers=atty).status_code == 422
    listing = client.get("/api/matters", headers=para).json()  # the whole firm sees all matters
    assert [m["matter_number"] for m in listing] == ["M-0002", "M-0001"]
    assert listing[0]["document_count"] == 0 and listing[0]["last_run_status"] is None
    assert client.get("/api/matters/999", headers=atty).status_code == 404


def test_document_upload_redacts_and_stores_no_raw_pii(client, atty, session_factory, settings):
    mid = make_matter(client, atty)
    r = upload(client, atty, mid)
    assert r.status_code == 201, r.text
    doc = r.json()
    assert "123-45-6789" not in doc["redacted_text"] and "jane.doe@example.com" not in doc["redacted_text"]
    types = {e["entity_type"] for e in doc["entities"]}
    assert {"US_SSN", "EMAIL_ADDRESS"} <= types and doc["redaction_count"] == len(doc["entities"])
    with session_factory() as s:
        row = s.get(Document, doc["id"])
        blob = json.dumps([row.redacted_text, row.redaction_entities])
        assert "123-45-6789" not in blob and "jane.doe" not in blob
        audit_rows = s.query(AuditLog).filter_by(action="document.uploaded").all()
        assert len(audit_rows) == 1 and "123-45-6789" not in json.dumps(audit_rows[0].detail)
        assert audit_rows[0].detail["redactions"] == doc["redaction_count"]
    assert not list(settings.upload_dir.rglob("*.txt"))  # raw bytes never written to disk
    detail = client.get(f"/api/matters/{mid}", headers=atty).json()
    assert detail["document_count"] == 1 and detail["documents"][0]["id"] == doc["id"]


def test_upload_rejections(client, atty, monkeypatch):
    mid = make_matter(client, atty)
    r = client.post(f"/api/matters/{mid}/documents", headers=atty, files={"file": ("x.exe", b"MZ\x00", "application/x")})
    assert r.status_code == 422
    assert client.post("/api/matters/999/documents", headers=atty,
                       files={"file": ("a.txt", b"hello world text", "text/plain")}).status_code == 404
    monkeypatch.setattr("stare.api.routers.documents.MAX_UPLOAD_BYTES", 10)
    assert upload(client, atty, mid, "x" * 50).status_code == 413


def test_redaction_correction_is_audited_and_reredacted(client, atty, session_factory):
    mid = make_matter(client, atty)
    doc = upload(client, atty, mid).json()
    r = client.patch(f"/api/documents/{doc['id']}", headers=atty,
                     json={"redacted_text": "Corrected text. Call me at jane@leak.example.com"})
    assert r.status_code == 200
    assert "jane@leak.example.com" not in r.json()["redacted_text"]
    with session_factory() as s:
        assert s.query(AuditLog).filter_by(action="document.redaction_edited").count() == 1


def test_redaction_correction_keeps_placeholders_distinct_and_counted(client, atty):
    mid = make_matter(client, atty)
    doc = upload(client, atty, mid).json()
    assert doc["redaction_count"] >= 1
    edited = doc["redacted_text"] + " Her supervisor John Smith saw it; call 972-555-0100."
    out = client.patch(f"/api/documents/{doc['id']}", headers=atty, json={"redacted_text": edited}).json()
    assert "John Smith" not in out["redacted_text"] and "972-555-0100" not in out["redacted_text"]
    assert out["redaction_count"] == doc["redaction_count"] + 2  # earlier redactions still count
    new = [e for e in out["entities"] if e["start"] >= len(doc["redacted_text"])]
    old = {e["placeholder"] for e in doc["entities"]}
    assert new and not ({e["placeholder"] for e in new} & old)  # a new person never reuses [PERSON_1]
    for e in out["entities"]:
        assert out["redacted_text"][e["start"]:e["end"]] == e["placeholder"]


# ----------------------------------------------------------------------------------------------- runs


def start_to_plan_gate(client, headers):
    mid = make_matter(client, headers)
    assert upload(client, headers, mid).status_code == 201
    r = client.post(f"/api/matters/{mid}/runs", headers=headers)
    assert r.status_code == 201, r.text
    run = r.json()
    assert run["status"] == "awaiting_plan_approval" and run["plan"]["items"]
    return mid, run


def test_start_run_requires_documents(client, atty):
    mid = make_matter(client, atty)
    assert client.post(f"/api/matters/{mid}/runs", headers=atty).status_code == 409


def test_full_run_roles_and_export(client, atty, para, admin, session_factory):
    mid, run = start_to_plan_gate(client, para)  # a paralegal may start runs
    rid = run["id"]
    # paralegals cannot approve or export
    assert client.post(f"/api/runs/{rid}/plan/approve", json=run["plan"], headers=para).status_code == 403
    assert client.get(f"/api/runs/{rid}/export.docx", headers=para).status_code == 403
    assert client.post(f"/api/runs/{rid}/cancel", headers=para).status_code == 403
    # but may edit the case record
    rec = {**run["case_record"], "procedural_posture": "Pre-suit"}
    assert client.put(f"/api/runs/{rid}/case-record", json=rec, headers=para).status_code == 200

    # memo gate cannot be skipped
    assert client.post(f"/api/runs/{rid}/memo/approve", headers=atty).status_code == 409
    r = client.post(f"/api/runs/{rid}/plan/approve", json=run["plan"], headers=atty)
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "awaiting_memo_approval" and r.json()["memo"]
    assert client.get(f"/api/runs/{rid}/export.docx", headers=atty).status_code == 409  # not approved yet

    # an edited memo with a hallucinated citation is rejected with problems
    bad = dict(r.json()["memo"])
    bad["rule"] += " See Varghese v. China Southern Airlines, 925 F.3d 1339 (11th Cir. 2019)."
    rr = client.post(f"/api/runs/{rid}/memo/approve", json=bad, headers=atty)
    assert rr.status_code == 422 and rr.json()["problems"]

    assert client.post(f"/api/runs/{rid}/memo/approve", headers=atty).json()["status"] == "approved"
    exp = client.get(f"/api/runs/{rid}/export.docx", headers=atty)
    assert exp.status_code == 200
    assert exp.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert f"STARE-memo-M-0001-run{rid}.docx" in exp.headers["content-disposition"]
    text = "\n".join(p.text for p in DocxDocument(io.BytesIO(exp.content)).paragraphs)
    assert "Rule" in text
    assert client.get(f"/api/runs/{rid}", headers=para).json()["status"] == "exported"
    assert client.get("/api/runs/999", headers=atty).status_code == 404

    audit = client.get("/api/audit", headers=admin).json()
    actions = [a["action"] for a in audit]
    for expected in ("auth.login", "matter.created", "document.uploaded", "run.started", "plan.approved",
                     "memo.approved", "memo.exported"):
        assert expected in actions
    assert audit[0]["id"] > audit[-1]["id"]  # newest first
    only = client.get(f"/api/audit?run_id={rid}&action=memo.exported", headers=admin).json()
    assert [a["action"] for a in only] == ["memo.exported"]
    assert client.get("/api/audit", headers=atty).status_code == 403


def test_cancel_and_changes(client, atty):
    _, run = start_to_plan_gate(client, atty)
    rid = run["id"]
    client.post(f"/api/runs/{rid}/plan/approve", json=run["plan"], headers=atty)
    r = client.post(f"/api/runs/{rid}/memo/changes", json={"instruction": "More on causation"}, headers=atty)
    assert r.status_code == 200 and r.json()["memo_version"] == 2
    assert client.post(f"/api/runs/{rid}/cancel", headers=atty).json()["status"] == "cancelled"
    assert client.post(f"/api/runs/{rid}/plan/approve", json=run["plan"], headers=atty).status_code == 409


def test_events_sse_replays_steps_and_status(client, atty, service):
    _, run = start_to_plan_gate(client, atty)
    rid = run["id"]
    client.post(f"/api/runs/{rid}/cancel", headers=atty)  # terminal: the stream replays then closes
    token = atty["Authorization"].split()[1]
    assert client.get(f"/api/runs/{rid}/events").status_code == 401
    events, current = [], {}
    with client.stream("GET", f"/api/runs/{rid}/events?token={token}") as resp:
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("text/event-stream")
        for line in resp.iter_lines():
            if line.startswith("event:"):
                current = {"event": line.split(":", 1)[1].strip()}
            elif line.startswith("data:"):
                current["data"] = json.loads(line.split(":", 1)[1])
                events.append(current)
    traces = [e for e in events if e["event"] == "trace"]
    assert traces and traces[0]["data"]["agent"] == "intake"
    assert [t["data"]["seq"] for t in traces] == sorted(t["data"]["seq"] for t in traces)
    assert events[-1] == {"event": "status", "data": {"run_id": rid, "status": "cancelled", "error": None}}
    assert not service.bus._subs.get(rid)  # subscription released


def test_graph_and_opinion(client, atty, session_factory):
    _, run = start_to_plan_gate(client, atty)
    rid = run["id"]
    client.post(f"/api/runs/{rid}/plan/approve", json=run["plan"], headers=atty)
    with session_factory() as s:
        for oid, name, cite in ((1, "McDonnell Douglas Corp. v. Green", "411 U.S. 792"),
                                (2, "Nassar", "570 U.S. 338")):
            s.add(Opinion(id=oid, court_id=None, case_name=name, citation=cite, plain_text=f"text of {name}",
                          source="courtlistener", source_id=str(100 + oid)))
        s.flush()
        s.add(CitationEdge(citing_id=2, cited_id=1, treatment="positive"))
        s.commit()
    g = client.get(f"/api/runs/{rid}/graph", headers=atty).json()
    ids = {n["id"] for n in g["nodes"]}
    assert {"op:1", "op:2"} <= ids and any(n["adverse"] for n in g["nodes"])
    assert {"source": "op:2", "target": "op:1", "treatment": "positive"} in g["edges"]
    op = client.get("/api/opinions/1", headers=atty).json()
    assert op["text"] == "text of McDonnell Douglas Corp. v. Green" and op["url"].endswith("/opinion/101/")
    assert client.get("/api/opinions/77", headers=atty).status_code == 404


# ----------------------------------------------------------------------------------------------- firewall


MATA = ("Plaintiff relies on Varghese v. China Southern Airlines Co., 925 F.3d 1339 (11th Cir. 2019), "
        "which holds that the statute of limitations is tolled. The court agreed.")


def test_firewall_check_real_firewall_blocks_fabricated_citation(client, atty, service, session_factory, settings):
    assert settings.courtlistener_enabled is False
    service.deps.__dict__["firewall"] = CitationFirewall(session_factory, settings)
    r = client.post("/api/firewall/check", json={"text": MATA}, headers=atty)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["checks"] and all(c["status"] in ("blocked", "unresolved") for c in body["checks"])
    assert body["exportable"] is False and "unresolved" in body["summary"] + "blocked"
    assert "925 F.3d 1339" not in body["cleaned_text"] and body["removed"]
    with session_factory() as s:
        a = s.query(AuditLog).filter_by(action="firewall.checked").one()
        assert a.detail["citations"] == len(body["checks"]) and "Varghese" not in json.dumps(a.detail)
    assert client.post("/api/firewall/check", json={"text": MATA}).status_code == 401


def test_firewall_check_redacts_pii_first(client, atty):
    r = client.post("/api/firewall/check", headers=atty,
                    json={"text": "Client SSN 123-45-6789 sued under Burlington, 548 U.S. 53 (2006)."})
    body = r.json()
    assert "123-45-6789" not in body["cleaned_text"] and body["redaction_count"] >= 1
    assert body["counts"]["verified"] == 1 and body["exportable"] is True


# ----------------------------------------------------------------------------------------------- settings


def test_settings_admin_only(client, admin, atty, session_factory):
    s = client.get("/api/settings", headers=admin).json()
    assert s["retention_days"] == 365 and s["llm_provider"] == "offline"
    assert client.get("/api/settings", headers=atty).status_code == 403
    assert client.patch("/api/settings", json={"retention_days": 90}, headers=admin).json()["retention_days"] == 90
    assert client.get("/api/settings", headers=admin).json()["retention_days"] == 90
    assert client.patch("/api/settings", json={"retention_days": 0}, headers=admin).status_code == 422
    with session_factory() as s_:
        assert s_.query(AuditLog).filter_by(action="settings.updated").count() == 1


def test_case_record_edit_redrafts_plan_and_redacts_typed_text(client, atty, session_factory):
    _, run = start_to_plan_gate(client, atty)
    rid = run["id"]
    statutes = {i.get("statute") for i in run["plan"]["items"]}
    assert "ada" not in statutes
    rec = {**run["case_record"],
           "claims": run["case_record"]["claims"] + [{"statute": "ada", "theory": "failure_to_accommodate", "description": ""}],
           "facts_summary": "Plaintiff was fired. Reach her at jane@leak.example.com or 214-555-0199.",
           "parties": run["case_record"]["parties"] + [{"role": "other", "name": "Mrs. Karen Miller"}]}
    r = client.put(f"/api/runs/{rid}/case-record", json=rec, headers=atty)
    assert r.status_code == 200, r.text
    out = r.json()
    assert "ada" in {i.get("statute") for i in out["plan"]["items"]}  # the new claim is in the plan
    blob = json.dumps(out["case_record"])
    assert "jane@leak.example.com" not in blob and "214-555-0199" not in blob and "Karen Miller" not in blob
    assert out["case_record"]["redaction_count"] >= run["case_record"]["redaction_count"] + 3
    kept = {p["name"] for p in run["case_record"]["parties"]}
    assert kept <= {p["name"] for p in out["case_record"]["parties"]}  # existing names are not re-redacted
    # the approved plan is the redrafted one
    r = client.post(f"/api/runs/{rid}/plan/approve", json=out["plan"], headers=atty)
    assert r.status_code == 200 and r.json()["status"] == "awaiting_memo_approval"


def test_memo_change_instruction_is_redacted(client, atty):
    _, run = start_to_plan_gate(client, atty)
    rid = run["id"]
    client.post(f"/api/runs/{rid}/plan/approve", json=run["plan"], headers=atty)
    r = client.post(f"/api/runs/{rid}/memo/changes", json={"instruction": "Her SSN is 123-45-6789; add causation."},
                    headers=atty)
    assert r.status_code == 200 and "123-45-6789" not in json.dumps(r.json()["memo"])


def test_plan_text_typed_by_the_attorney_is_redacted(client, atty):
    _, run = start_to_plan_gate(client, atty)
    rid = run["id"]
    plan = json.loads(json.dumps(run["plan"]))
    drafted = plan["items"][0]["elements"][0]
    plan["instructions"] = "Focus on retaliation; client SSN 123-45-6789, email jane@leak.example.com."
    plan["items"][0]["elements"].append("Mrs. Karen Miller complained to HR, call 214-555-0199")
    r = client.post(f"/api/runs/{rid}/plan/approve", json=plan, headers=atty)
    assert r.status_code == 200, r.text
    blob = json.dumps(r.json()["plan"]) + json.dumps(r.json()["memo"])
    for leaked in ("123-45-6789", "jane@leak.example.com", "Karen Miller", "214-555-0199"):
        assert leaked not in blob
    assert drafted in r.json()["plan"]["items"][0]["elements"]  # drafted text is not re-redacted
    assert "Focus on retaliation" in r.json()["plan"]["instructions"]
