"""Editable Word export of an approved IRAC memo (FR-12, US-05)."""

from __future__ import annotations

import io
import re
from datetime import UTC, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from stare import DISCLAIMER
from stare.schemas import CitationCheck, Memo

_STATUS_LABEL = {"verified": "VERIFIED", "warning": "WARNING", "blocked": "BLOCKED", "unresolved": "UNRESOLVED"}
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


class MemoExporter:
    def to_docx(self, memo: Memo, run_meta: dict) -> bytes:
        """run_meta keys (all optional): matter_number, matter_title, court, prepared_by, approved_by,
        approved_at, run_id, checks (list of CitationCheck dicts)."""
        doc = Document()
        self._styles(doc)

        title = doc.add_heading(memo.title or "Research Memorandum", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self._header_table(doc, memo, run_meta)

        self._section(doc, "Question Presented", memo.question_presented)
        self._section(doc, "Issue", memo.issue)
        self._section(doc, "Rule", memo.rule)
        self._section(doc, "Application", memo.application)
        self._section(doc, "Conclusion", memo.conclusion)
        self._section(doc, "Adverse Authority (Opposing Counsel review)", memo.adverse_authority or "None identified.")

        if memo.warnings:
            doc.add_heading("Citation Firewall Warnings", level=1)
            for w in memo.warnings:
                doc.add_paragraph(w, style="List Bullet")

        doc.add_heading("Authorities Cited", level=1)
        for c in memo.citations:
            p = doc.add_paragraph()
            p.add_run(f"{c.marker} ").bold = True
            if c.case_name:
                p.add_run(c.case_name).italic = True
                p.add_run(", ")
            if c.url:
                self._hyperlink(p, c.citation, c.url)
            else:
                p.add_run(c.citation)
            p.add_run(f"  [{_STATUS_LABEL.get(c.status, c.status.upper())}]").font.size = Pt(8)

        checks = [CitationCheck.model_validate(c) for c in run_meta.get("checks") or []]
        if checks or memo.removed_citations:
            self._verification_appendix(doc, checks, memo.removed_citations)

        self._footer(doc, memo.disclaimer or DISCLAIMER)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------------------------------

    def _styles(self, doc) -> None:
        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)

    def _header_table(self, doc, memo: Memo, meta: dict) -> None:
        approved_at = meta.get("approved_at")
        if isinstance(approved_at, datetime):
            approved_at = approved_at.strftime("%d %B %Y %H:%M UTC")
        rows = [
            ("Matter", " ".join(x for x in (meta.get("matter_number"), meta.get("matter_title")) if x)),
            ("Court", meta.get("court") or ""),
            ("Prepared by", meta.get("prepared_by") or "STARE research agents"),
            ("Approved by", meta.get("approved_by") or ""),
            ("Approved at", approved_at or ""),
            ("Run", str(meta.get("run_id") or "")),
            ("Exported", datetime.now(UTC).strftime("%d %B %Y %H:%M UTC")),
            ("Drafted with", memo.generated_by),
        ]
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for k, v in rows:
            if not v:
                continue
            cells = table.add_row().cells
            cells[0].text = k
            cells[1].text = v
        doc.add_paragraph()

    def _section(self, doc, heading: str, body: str) -> None:
        doc.add_heading(heading, level=1)
        for block in re.split(r"\n\s*\n", (body or "").strip()):
            if not block.strip():
                continue
            paragraph: list[str] = []
            for ln in block.strip().splitlines():
                if re.match(r"^\s*[-*•]\s+", ln):
                    if paragraph:
                        self._rich(doc.add_paragraph(), " ".join(paragraph))
                        paragraph = []
                    self._rich(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", ln))
                else:
                    paragraph.append(ln.strip())
            if paragraph:
                self._rich(doc.add_paragraph(), " ".join(paragraph))

    def _rich(self, p, text: str) -> None:
        for part in _INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("*") and part.endswith("*"):
                p.add_run(part[1:-1]).italic = True
            else:
                p.add_run(part)

    def _verification_appendix(self, doc, checks: list[CitationCheck], removed: list[str]) -> None:
        doc.add_page_break()
        doc.add_heading("Appendix: Citation Verification Record", level=1)
        doc.add_paragraph(
            "Every citation proposed by the research agents was checked by the Citation Firewall: "
            "Gate 1 extraction, Gate 2 existence in court records, Gate 3 quotation match and Gate 4 "
            "later negative treatment. Blocked citations were removed before the memo was drafted."
        )
        if checks:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            for i, h in enumerate(("Citation", "Status", "Gates", "Reason")):
                hdr[i].text = h
            for c in checks:
                cells = table.add_row().cells
                cells[0].text = c.citation.raw or c.citation.cite
                cells[1].text = _STATUS_LABEL.get(c.status, c.status)
                cells[2].text = " ".join(f"G{g.gate}:{g.outcome}" for g in c.gates)
                cells[3].text = c.reason
        if removed:
            doc.add_paragraph()
            doc.add_paragraph("Citations removed by the Firewall:").runs[0].bold = True
            for r in removed:
                doc.add_paragraph(r, style="List Bullet")

    def _footer(self, doc, text: str) -> None:
        for section in doc.sections:
            p = section.footer.paragraphs[0]
            p.text = text
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    def _hyperlink(self, paragraph, text: str, url: str) -> None:
        r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1F4E9A")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rpr.append(color)
        rpr.append(underline)
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        link.append(run)
        paragraph._p.append(link)
