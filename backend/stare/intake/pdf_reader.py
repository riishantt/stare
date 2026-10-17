"""Document text extraction: PDF (PyMuPDF, Tesseract OCR fallback), DOCX and TXT."""

from __future__ import annotations

import io
import re
import shutil
from pathlib import Path

from stare.schemas import ExtractedDocument, PageText

MIN_PAGE_CHARS = 40  # below this a PDF page is treated as scanned and OCRed
OCR_DPI = 250
_TESSERACT_FALLBACK = "/opt/homebrew/bin/tesseract"


class DocumentReadError(ValueError):
    """The file is unsupported, corrupt, encrypted or could not be read."""


def clean_text(text: str) -> str:
    """De-hyphenate line-break hyphens, normalise whitespace; keep line and paragraph breaks
    (caption layout matters to the extractor)."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x0c", "\n").replace("\xa0", " ")
    text = text.replace("­", "")
    text = re.sub(r"(?<=[A-Za-z])-[ \t]*\n[ \t]*(?=[a-z])", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _ocr_image(png: bytes) -> str:
    import pytesseract
    from PIL import Image

    if not shutil.which("tesseract") and Path(_TESSERACT_FALLBACK).exists():
        pytesseract.pytesseract.tesseract_cmd = _TESSERACT_FALLBACK
    try:
        with Image.open(io.BytesIO(png)) as img:
            return pytesseract.image_to_string(img.convert("L"), config="--psm 4")
    except pytesseract.TesseractNotFoundError as exc:
        raise DocumentReadError("OCR is required for this scanned PDF but Tesseract is not installed") from exc


def _read_pdf(data: bytes) -> list[PageText]:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise DocumentReadError("The file could not be opened as a PDF (corrupt or not a PDF)") from exc
    pages: list[PageText] = []
    try:
        if doc.needs_pass:
            raise DocumentReadError("The PDF is password-protected")
        for i, page in enumerate(doc, start=1):
            raw = page.get_text("text") or ""
            ocr = False
            if len(raw.strip()) < MIN_PAGE_CHARS:
                pix = page.get_pixmap(dpi=OCR_DPI, colorspace=pymupdf.csGRAY)
                ocr_text = _ocr_image(pix.tobytes("png"))
                del pix
                if len(ocr_text.strip()) > len(raw.strip()):
                    raw, ocr = ocr_text, True
            pages.append(PageText(page_no=i, text=clean_text(raw), ocr_used=ocr))
    except DocumentReadError:
        raise
    except Exception as exc:
        raise DocumentReadError(f"Failed while reading the PDF: {exc}") from exc
    finally:
        doc.close()
    return pages


def _read_docx(data: bytes) -> list[PageText]:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentReadError("The file could not be opened as a .docx document") from exc
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [PageText(page_no=1, text=clean_text("\n".join(parts)))]


def _read_txt(data: bytes) -> list[PageText]:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    else:
        raise DocumentReadError("The text file has an unsupported encoding")
    if "\x00" in text:
        raise DocumentReadError("The file does not look like a text document")
    return [PageText(page_no=1, text=clean_text(text))]


def extract_document(data: bytes, file_name: str) -> ExtractedDocument:
    if not data:
        raise DocumentReadError("The file is empty")
    ext = Path(file_name).suffix.lower()
    if ext == ".pdf" or (not ext and data[:5] == b"%PDF-"):
        pages = _read_pdf(data)
    elif ext == ".docx":
        pages = _read_docx(data)
    elif ext in (".txt", ".text"):
        pages = _read_txt(data)
    else:
        raise DocumentReadError(f"Unsupported file type '{ext or file_name}': upload a PDF, DOCX or TXT file")
    text = "\n\n".join(p.text for p in pages if p.text)
    if not text.strip():
        raise DocumentReadError("No readable text was found in the document")
    return ExtractedDocument(
        text=text, pages=pages, ocr_used=any(p.ocr_used for p in pages), page_count=len(pages)
    )
